"""Document text extraction.

The pipeline talks to `DocumentExtractor`. v1 ships pure-Python extractors for
PDF, DOCX and plain text; Azure AI Document Intelligence Layout would be another
implementation of the same protocol returning the same `ExtractedDocument`.

What the local extractors give up relative to Document Intelligence: table
structure, reading order on multi-column layouts, and any OCR at all - a scanned
PDF extracts as empty and is rejected rather than silently indexed as a blank
document. That trade is recorded in docs/decisions.md entry 14.

The unit of extraction is a **block**, not a line: a block carries its page
number and whether it looked like a heading, which is what makes section-aware
chunking possible downstream.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Protocol

from app.core.logging import get_logger

logger = get_logger(__name__)

PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
TEXT_TYPES = {"text/plain", "text/markdown", "text/x-markdown", "application/octet-stream"}

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class ExtractionError(Exception):
    """The file could not be turned into text. Carries a user-safe reason."""


@dataclass(slots=True)
class Block:
    """One paragraph-ish run of text, with where it came from."""

    text: str
    page_number: int
    # Set by an extractor that knows structurally (DOCX heading styles); left
    # None for formats where it has to be inferred later from the text itself.
    is_heading: bool | None = None


@dataclass(slots=True)
class ExtractedDocument:
    blocks: list[Block] = field(default_factory=list)
    page_count: int = 0
    # Extractor-specific notes - e.g. "encrypted", "no extractable text".
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(block.text for block in self.blocks)

    @property
    def is_empty(self) -> bool:
        return not any(block.text.strip() for block in self.blocks)


class DocumentExtractor(Protocol):
    def supports(self, *, content_type: str, filename: str) -> bool: ...

    def extract(self, data: bytes, *, filename: str) -> ExtractedDocument: ...


class PlainTextExtractor:
    """`.txt` and `.md`. One block per blank-line-separated paragraph."""

    def supports(self, *, content_type: str, filename: str) -> bool:
        return content_type in TEXT_TYPES or filename.lower().endswith((".txt", ".md"))

    def extract(self, data: bytes, *, filename: str) -> ExtractedDocument:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            # Clinical exports out of older systems are often cp1252.
            text = data.decode("cp1252", errors="replace")

        blocks = [
            Block(text=paragraph.strip(), page_number=1)
            for paragraph in text.split("\n\n")
            if paragraph.strip()
        ]
        return ExtractedDocument(blocks=blocks, page_count=1)


class PdfExtractor:
    """`.pdf` via pypdf. Text layer only - no OCR."""

    def supports(self, *, content_type: str, filename: str) -> bool:
        return content_type in PDF_TYPES or filename.lower().endswith(".pdf")

    def extract(self, data: bytes, *, filename: str) -> ExtractedDocument:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
        except PdfReadError as exc:
            raise ExtractionError(f"The PDF could not be read: {exc}") from exc

        if reader.is_encrypted:
            # An empty-password PDF is common and decrypts silently.
            try:
                if reader.decrypt("") == 0:
                    raise ExtractionError("The PDF is password protected.")
            except NotImplementedError as exc:
                raise ExtractionError("The PDF uses an unsupported encryption scheme.") from exc

        document = ExtractedDocument(page_count=len(reader.pages))
        for index, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001 - one bad page must not lose the rest
                logger.warning("pdf_page_extraction_failed", page=index, error=str(exc))
                document.warnings.append(f"Page {index} could not be read.")
                continue
            for paragraph in page_text.split("\n\n"):
                if paragraph.strip():
                    document.blocks.append(Block(text=paragraph.strip(), page_number=index))

        if document.is_empty:
            raise ExtractionError(
                "No extractable text. The PDF is most likely a scan; OCR is not part of v1."
            )
        return document


class DocxExtractor:
    """`.docx` via python-docx.

    The one format where headings are known rather than guessed: Word records
    the paragraph style, so `is_heading` comes from the document itself.
    """

    def supports(self, *, content_type: str, filename: str) -> bool:
        return content_type in DOCX_TYPES or filename.lower().endswith(".docx")

    def extract(self, data: bytes, *, filename: str) -> ExtractedDocument:
        import docx
        from docx.opc.exceptions import PackageNotFoundError

        try:
            document_file = docx.Document(io.BytesIO(data))
        except PackageNotFoundError as exc:
            raise ExtractionError("The file is not a valid .docx package.") from exc

        extracted = ExtractedDocument(page_count=1)
        for paragraph in document_file.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            extracted.blocks.append(
                Block(text=text, page_number=1, is_heading=style.startswith("Heading"))
            )

        # Tables carry lab values and medication lists; losing them would lose
        # the most citable content in a clinical document.
        for table in document_file.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    extracted.blocks.append(
                        Block(text=" | ".join(cells), page_number=1, is_heading=False)
                    )

        if extracted.is_empty:
            raise ExtractionError("The document contains no text.")
        return extracted


class ExtractorRegistry:
    """Picks the first extractor that claims the file."""

    def __init__(self, extractors: list[DocumentExtractor] | None = None) -> None:
        self._extractors: list[DocumentExtractor] = extractors or [
            PdfExtractor(),
            DocxExtractor(),
            PlainTextExtractor(),
        ]

    def for_file(self, *, content_type: str, filename: str) -> DocumentExtractor:
        for extractor in self._extractors:
            if extractor.supports(content_type=content_type, filename=filename):
                return extractor
        raise ExtractionError(
            f"Unsupported file type '{content_type or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    def extract(self, data: bytes, *, content_type: str, filename: str) -> ExtractedDocument:
        extractor = self.for_file(content_type=content_type, filename=filename)
        return extractor.extract(data, filename=filename)
